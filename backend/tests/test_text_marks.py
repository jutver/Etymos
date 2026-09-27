"""
Tests for extractor.compute_text_marks: inline bold/italic kept from the
original file as {start, end, style} ranges over the block text.

The invariant: every mark's range slices exactly the words that were styled
in the source, and a block whose text can't be reproduced gets no marks
rather than marks in the wrong place.
"""

from docx import Document

from extractor import clean_text, compute_text_marks
from extract_docx import _paragraph_block


def _styled(text, marks, style):
    return [text[m["start"]:m["end"]] for m in marks if m["style"] == style]


def test_pdf_spans_joined_with_spaces():
    pieces = [("We use", False, False), ("PhoBERT", True, False), ("for", False, False), ("hate speech", False, True)]
    text = clean_text(" ".join(p[0] for p in pieces))
    marks = compute_text_marks(pieces, " ", text)
    assert _styled(text, marks, "bold") == ["PhoBERT"]
    assert _styled(text, marks, "italic") == ["hate speech"]


def test_adjacent_bold_spans_become_one_mark():
    pieces = [("Hate", True, False), ("speech", True, False), ("(HATE)", True, False), ("is", False, False)]
    text = clean_text(" ".join(p[0] for p in pieces))
    marks = compute_text_marks(pieces, " ", text)
    assert _styled(text, marks, "bold") == ["Hate speech (HATE)"]


def test_list_marker_stripped_from_the_front_is_matched_as_suffix():
    pieces = [("1.", False, False), ("Warning:", True, False), ("offensive content", False, False)]
    marks = compute_text_marks(pieces, " ", "Warning: offensive content")
    assert _styled("Warning: offensive content", marks, "bold") == ["Warning:"]


def test_unreproducible_text_gets_no_marks():
    pieces = [("Completely", True, False), ("different", False, False)]
    assert compute_text_marks(pieces, " ", "Nothing alike here") == []


def test_plain_text_gets_no_marks():
    assert compute_text_marks([("plain", False, False), ("words", False, False)], " ", "plain words") == []


def test_docx_runs_keep_bold_and_italic():
    paragraph = Document().add_paragraph()
    paragraph.add_run("Mô hình ")
    paragraph.add_run("PhoBERT-CNN").bold = True
    paragraph.add_run(" đạt kết quả ")
    paragraph.add_run("tốt nhất").italic = True
    paragraph.add_run(".")

    block = _paragraph_block(paragraph, page=1)
    text = block["text"]
    assert text == "Mô hình PhoBERT-CNN đạt kết quả tốt nhất."
    assert _styled(text, block["marks"], "bold") == ["PhoBERT-CNN"]
    assert _styled(text, block["marks"], "italic") == ["tốt nhất"]


def test_heading_line_keeps_marks_when_shown_as_paragraph():
    from extractor import _enforce_heading_sequence, build_structured_blocks

    line = {
        "text": "1 Faculty of Information Science", "page": 1, "size": 14.0,
        "font": "Bold", "words": 5, "bold": False,
        "spans": [("1", False, False), ("Faculty", True, False), ("of Information Science", False, False)],
    }
    blocks = build_structured_blocks([line], body_size=10.0)
    assert blocks[0]["type"] == "heading"
    sections = {"introduction": [
        {"type": "heading", "level": 2, "list_type": None, "text": "1 Introduction", "page": 1},
        {**blocks[0], "level": 2},
    ]}
    _enforce_heading_sequence(sections, ["introduction"])
    demoted = sections["introduction"][1]
    assert demoted["type"] == "paragraph"
    assert _styled(demoted["text"], demoted["marks"], "bold") == ["Faculty"]
