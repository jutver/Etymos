"""
extract_docx.py
----------------
.docx counterpart to extractor.py's PDF extraction. Produces the same
structured-block shape extractor.extract_structured_sections_from_pdf()
does — {"title": str, "sections": {section_name: [block, ...]}}, where each
block is {"type", "level", "list_type", "text", "page", "rows"?} — so both
paths feed document_model.build_document() identically and the Document
view renders either kind of upload the same way.

Unlike the PDF path, .docx documents are essentially never organized as
abstract/introduction/method/... (that taxonomy is PDF-academic-paper
specific), so there is no reliable heading vocabulary to sort content by
section. Everything therefore lands in the single "body" catch-all section
that extractor.py already defines for exactly this "just keep everything,
in order, and let block-type carry the structure" situation — see
extractor.FALLBACK_SECTION.

Unlike PDF text extraction (which reads glyph positions off a rendered
page), python-docx reads the document's own XML structure directly, so
heading level and list membership come from real paragraph styles / list
formatting rather than a font-size heuristic. This is why the .docx path
recovers headings/lists more reliably than the PDF path can.
"""

from __future__ import annotations

import logging

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from extractor import clean_text, LIST_MARKER_RE, BULLET_MARKER_CHARS, FALLBACK_SECTION

logger = logging.getLogger(__name__)

HEADING_STYLE_LEVEL = {
    "title": 1,
    "heading 1": 1,
    "heading 2": 2,
    "heading 3": 3,
    "heading 4": 3,
    "heading 5": 3,
    "heading 6": 3,
}

# Word does not expose real page boundaries at the XML level — pagination
# is a rendering-time concern that depends on fonts, margins, and the
# renderer. We honour EXPLICIT page breaks (Ctrl+Enter / "Insert > Page
# Break") when present in the XML, and otherwise bucket every N blocks into
# a synthetic page so the Document view still renders real blank-space
# page gaps instead of one endless strip. This is a known, documented
# limitation of the .docx path: synthetic pages will not line up with
# Word's own pagination exactly.
BLOCKS_PER_SYNTHETIC_PAGE = 14


def _iter_block_items(document: Document):
    """
    Yield paragraphs and tables in true document order.

    python-docx's `document.paragraphs` / `document.tables` accessors each
    flatten their own kind and lose how they interleave — a table in the
    middle of the text would otherwise get moved to the end. Walking the
    body's XML children directly (the standard python-docx idiom for this)
    preserves the real reading order.
    """
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _has_hard_page_break(paragraph: Paragraph) -> bool:
    for br in paragraph._p.findall(".//" + qn("w:br")):
        if br.get(qn("w:type")) == "page":
            return True
    return False


def _list_type_from_style(paragraph: Paragraph) -> str | None:
    style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
    if "bullet" in style_name:
        return "bullet"
    if "number" in style_name:
        return "number"

    p_pr = paragraph._p.pPr
    if p_pr is not None and p_pr.numPr is not None:
        # Direct (non-styled) numbering is applied, but resolving whether
        # it's a bullet or an ordered numbering scheme requires walking
        # numbering.xml's abstractNum definitions. Default to "bullet" —
        # the more common case for ad-hoc list formatting — rather than
        # guessing wrong the other way; either way this is a display
        # detail, not a content-loss risk.
        return "bullet"

    return None


def _strip_manual_marker(text: str) -> tuple[str, str | None]:
    """Some documents fake a list with literal '- ' / '1. ' text instead of
    real list formatting. Recognize that the same way the PDF path does."""
    marker = LIST_MARKER_RE.match(text)
    if not marker:
        return text, None
    stripped = text[marker.end():].strip()
    if not stripped:
        return text, None
    marker_text = marker.group(0).strip()
    list_type = "bullet" if marker_text[:1] in BULLET_MARKER_CHARS else "number"
    return stripped, list_type


def _paragraph_block(paragraph: Paragraph, page: int) -> dict | None:
    # clean_text() already collapses any manual-line-break "\n" python-docx
    # renders inside paragraph.text; the extra .replace is a defensive
    # no-op belt-and-braces guard, not load-bearing today. See the comment
    # in _table_block for why a stray "\n" here would be a silent
    # structure-loss bug for the whole section, not just a cosmetic one.
    text = clean_text(paragraph.text).replace("\n", " ").strip()
    if not text:
        return None

    style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
    level = HEADING_STYLE_LEVEL.get(style_name)
    if level is not None:
        return {"type": "heading", "level": level, "list_type": None, "text": text, "page": page}

    list_type = _list_type_from_style(paragraph)
    if list_type:
        stripped, manual_type = _strip_manual_marker(text)
        return {
            "type": "list_item", "level": None,
            "list_type": manual_type or list_type,
            "text": stripped, "page": page,
        }

    stripped, manual_type = _strip_manual_marker(text)
    if manual_type:
        return {"type": "list_item", "level": None, "list_type": manual_type, "text": stripped, "page": page}

    return {"type": "paragraph", "level": None, "list_type": None, "text": text, "page": page}


# Maps the OOXML relationship part's content-type to a file extension for
# the Storage upload key (see report_store._extract_and_upload_images).
# Anything not listed falls back to "png" — a wrong extension on a rare
# format is a cosmetic filename mismatch, not a data-loss risk, since the
# bytes themselves (and the upload's explicit content-type) are unaffected.
_IMAGE_EXT_BY_CONTENT_TYPE = {
    "image/png": "png",
    "image/jpeg": "jpg",
    "image/jpg": "jpg",
    "image/gif": "gif",
    "image/bmp": "bmp",
    "image/tiff": "tiff",
    "image/x-emf": "emf",
    "image/x-wmf": "wmf",
}


def _paragraph_images(paragraph: Paragraph, document: Document, page: int) -> list[dict]:
    """
    Extract every inline picture embedded in `paragraph` as best-effort
    "image" blocks carrying raw bytes (`_image_bytes`/`_image_ext`) for a
    later upload step (report_store._extract_and_upload_images) to turn
    into a Storage `image_url`. These two keys are an in-process-only
    carrier — never JSON-serialized, always stripped before a report is
    persisted or returned.

    python-docx exposes an inline picture as a `<w:drawing>` run child
    containing a `<pic:pic>` element whose `<a:blip r:embed="...">`
    attribute is the relationship id; `document.part.related_parts[rId]`
    resolves that id to the actual image part or raises KeyError for a
    dangling relationship (rare, but not fatal here).

    Never raises: any failure extracting one picture is logged and
    skipped, exactly like the rest of this module's "one bad element
    should never drop the whole document" posture — image extraction is
    additive and must never break text extraction.
    """
    blocks: list[dict] = []
    try:
        pics = paragraph._p.xpath(".//pic:pic")
    except Exception:
        logger.exception("Failed to look up inline pictures in a docx paragraph.")
        return blocks

    seen_rids: set[str] = set()
    for pic in pics:
        try:
            for blip in pic.xpath(".//a:blip"):
                rid = blip.get(qn("r:embed"))
                if not rid or rid in seen_rids:
                    continue
                seen_rids.add(rid)

                part = document.part.related_parts.get(rid)
                if part is None:
                    continue
                image_bytes = part.blob
                if not image_bytes:
                    continue

                ext = _IMAGE_EXT_BY_CONTENT_TYPE.get(part.content_type, "png")
                blocks.append({
                    "type": "image", "level": None, "list_type": None,
                    "text": "", "page": page,
                    "_image_bytes": image_bytes, "_image_ext": ext,
                })
        except Exception:
            logger.exception("Failed to extract an inline docx image; skipping it.")
            continue

    return blocks


def _table_block(table: Table, page: int) -> dict | None:
    rows = []
    for row in table.rows:
        cells = [clean_text(cell.text) for cell in row.cells]
        if any(cells):
            rows.append(cells)
    if not rows:
        return None
    return {
        "type": "table", "level": None, "list_type": None,
        # No literal "\n" here on purpose: document_model._split_paragraphs
        # treats ANY newline as a block boundary (not just blank-line "\n\n"
        # runs), so a table whose flattened text contained one would split
        # into extra spans and desync the 1:1 block<->span correspondence
        # build_document() relies on to attach type/level/page metadata —
        # silently losing structure for the entire section, not just this
        # block. Rows are still available structured in "rows" below; this
        # "text" is only the matching-pipeline's plain-text view of it.
        "text": " ¶ ".join(" | ".join(r) for r in rows),
        "page": page,
        "rows": rows,
    }


def extract_structured_sections_from_docx(docx_path: str) -> dict:
    """
    Returns {"title": str, "sections": {"body": [block, ...]}} — the same
    shape extractor.extract_structured_sections_from_pdf() returns, minus
    the academic section buckets (see module docstring for why).
    """
    document = Document(docx_path)

    blocks: list[dict] = []
    title = ""
    page = 1
    blocks_on_page = 0

    for item in _iter_block_items(document):
        if isinstance(item, Table):
            block = _table_block(item, page)
            if block:
                blocks.append(block)
                blocks_on_page += 1
            continue

        style_name = (item.style.name or "").lower() if item.style else ""
        if not title and style_name == "title":
            candidate = clean_text(item.text)
            if candidate:
                title = candidate
                continue

        block = _paragraph_block(item, page)
        if block:
            blocks.append(block)
            blocks_on_page += 1

        # Best-effort inline-image extraction. Runs regardless of whether
        # the paragraph also had text (a figure often sits alone in its
        # own paragraph, in which case `block` above is None) — appended
        # right after the paragraph's own text block so an image lands
        # next to whatever caption/text shares its paragraph, matching
        # document order closely enough for the Document view.
        try:
            image_blocks = _paragraph_images(item, document, page)
        except Exception:
            image_blocks = []
            logger.exception("Image extraction failed for a docx paragraph; continuing without its images.")
        if image_blocks:
            blocks.extend(image_blocks)
            blocks_on_page += len(image_blocks)

        if _has_hard_page_break(item):
            page += 1
            blocks_on_page = 0
        elif blocks_on_page >= BLOCKS_PER_SYNTHETIC_PAGE:
            page += 1
            blocks_on_page = 0

    if not title:
        for b in blocks:
            if b["type"] == "heading":
                title = b["text"]
                break

    return {
        "title": title,
        "sections": {FALLBACK_SECTION: blocks} if blocks else {},
    }
